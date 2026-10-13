import json
import re
import sys
from pathlib import Path
from docx import Document

SOURCE_PATH = Path(sys.argv[1])
ROOT = Path(__file__).resolve().parents[1]
document = Document(SOURCE_PATH)

STOPWORDS = set("""
a an and are as at be because been being before but by can could did do does each either even for from
had has have he her here him his how i if in into is it its itself may might more most much must no nor
not now of on one only or other our ours out over she should so some such than that the their theirs them
themselves then there these they this those through to too under up us very was we were what when where
whether which while who whom whose why will with would you your yours first second third however therefore
also already later fairly quite really just still again together today yesterday tomorrow after behind many
notably economically politically socially during across year years time times day days way ways part parts thing things
""".split())

MONTHS = set("january february march april may june july august september october november december".split())

KNOWN_NAMES = set("""
azamgarh baron benjamin browne calcutta cambridge chunder constantinople curzon disraeli dutt edward george
hamilton hassan india john london macaulay mirza nathaniel nicholas paul reuter romesh russell seymour shirazi
st petersburg thomas victoria britain france russia persia egypt arabia hindustan tihrán tehran
""".split())

TITLES = set("""
baron emperor empress governor imam king khedive lord madam minister prince princes queen shah sir tsar
""".split())

ALLOWED_CAPITALIZED = set("""
British European Europeans English Indian Indians Persian Persians Ottoman Ottomans Muslim Muslims Hindu Hindus
Christian Christianity Shi’a Shi'a Raj Sepoy Social Darwinism
""".split())

THEME_WORDS = set("""
accurate attitudes authority backward banks biased bias boycott canal colonies colony competition concession
context control credible culture debt direct economic economy education effects empire evidence foreign government
imperial imperialism independence indirect industries industry influence investments labor lands limitation
limitations loans manufacture manufacturers manufacturing markets military missionaries monopoly motives
nationalism occupation oppression perspective political power primary profits propaganda property protest
protectorate protectorates rebellion reliable religion religious resources rights rule ruled ruler ruling secondary
social source spheres strategic superiority taxes territory trade value wealth
""".split())

ACTION_WORDS = set("""
argue argues attacked began believed borrowed bought called cancel cancelled changed compare connects controlled
created decide destroy destroyed divide divided drove encouraged explains forced gave govern governed imposed
include increased joined knew led lost made mention monopolized occupied paid planned presents promised protect
proves reduced reflect reported reveals ruled saw sell shared shows spoke stressed treated weakened wanted warned
worked wrote
""".split())

TOKEN_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ0-9]+(?:['’\-][A-Za-zÀ-ÖØ-öø-ÿ0-9]+)*")

def normalize(value):
    return value.lower().replace("’", "'")

def split_sentences(text):
    protected = re.sub(r"\b(?:Mr|Mrs|Ms|Dr|St|vs|etc)\.", lambda m: m.group(0).replace(".", "∯"), text)
    protected = re.sub(r"\b(?:[A-Z]\.){2,}", lambda m: m.group(0).replace(".", "∯"), protected)
    protected = re.sub(r"\b[A-Z]\.\s*(?=[A-Z][a-z])", lambda m: m.group(0).replace(".", "∯"), protected)
    parts = re.split(r"(?:(?<=[.!?][”’\"'])|(?<=[.!?])|(?<=[.!?]\)))\s+(?=[A-Z0-9“\"'])", protected)
    return [part.replace("∯", ".").strip() for part in parts if part.strip()]

def eligible(token, position):
    raw = token.group(0)
    word = normalize(raw)
    name_base = word[:-2] if word.endswith("'s") else word
    if not word or any(char.isdigit() for char in word) or word in STOPWORDS or word in MONTHS:
        return False
    if len(word) < 3 or name_base in KNOWN_NAMES or word in TITLES:
        return False
    if raw.isupper() and len(raw) <= 4:
        return False
    if raw[0].isupper() and raw not in ALLOWED_CAPITALIZED and position != 0:
        return False
    if position == 0 and name_base in KNOWN_NAMES | TITLES:
        return False
    return True

def word_score(word):
    normalized = normalize(word)
    score = 0
    if normalized in THEME_WORDS:
        score += 6
    if normalized in ACTION_WORDS or normalized.endswith(("ed", "ing", "ize", "ise", "fy")):
        score += 5
    if normalized.endswith(("tion", "sion", "ment", "ness", "ity", "ism", "ship", "ance", "ence")):
        score += 4
    if normalized.endswith(("al", "ic", "ive", "ous", "able", "ible", "ary", "ory")):
        score += 3
    score += min(len(normalized), 9) / 10
    return score

def choose_blank_words(sentence):
    matches = list(TOKEN_RE.finditer(sentence))
    candidates = []
    for position, match in enumerate(matches):
        if eligible(match, position):
            candidates.append((word_score(match.group(0)), position, match.group(0)))
    selected = sorted(sorted(candidates, key=lambda item: (-item[0], item[1]))[:10], key=lambda item: item[1])
    return [item[2] for item in selected]

def purpose_for(section, sentence_number, sentence_count):
    if section == "Introduction":
        labels = [
            "Identify the source, author, audience, and historical setting.",
            "Explain the source type or reporting context.",
            "State what the source reveals.",
            "Develop the main analytical claim.",
            "Give the overall credibility judgement.",
        ]
        return labels[min(sentence_number - 1, len(labels) - 1)]
    if section == "Body 1":
        return "Develop the content-and-context analysis with evidence and lesson concepts."
    if section == "Body 2":
        return "Evaluate how the source’s origin, audience, and purpose affect credibility."
    if section == "Body 3":
        return "Explain the source’s value and limitations for historians."
    return "Restate the argument and give a final credibility judgement."

def make_unit(letter):
    source_index = ord(letter) - ord("A")
    essay_table_index = 8 + source_index * 6
    lesson = "17-3" if letter in "ABCD" else "17-4"
    source_number = source_index + 1 if lesson == "17-3" else source_index - 3
    table = document.tables[essay_table_index]
    source_title = document.tables[essay_table_index - 4].cell(0, 0).text.strip()
    items = []
    item_id = 1
    for row in table.rows[1:]:
        heading_lines = [line.strip() for line in row.cells[0].text.splitlines() if line.strip()]
        section = heading_lines[0]
        focus = heading_lines[1] if len(heading_lines) > 1 else ""
        section_text = row.cells[1].text.strip()
        sentences = split_sentences(section_text)
        if " ".join(sentences) != section_text:
            raise ValueError(f"Sentence split changed source text for {letter} {section}")
        for section_sentence, sentence in enumerate(sentences, start=1):
            blank_words = choose_blank_words(sentence)
            if not blank_words:
                raise ValueError(f"No eligible blanks for {letter} {section} sentence {section_sentence}")
            items.append({
                "id": item_id,
                "essaySection": section,
                "sectionSentence": section_sentence,
                "sectionCount": len(sentences),
                "focus": focus,
                "term": section,
                "pos": f"Sentence {section_sentence} of {len(sentences)}",
                "definition": purpose_for(section, section_sentence, len(sentences)),
                "blankWords": blank_words,
                "prompt": sentence,
            })
            item_id += 1
    return {
        "code": f"{lesson}-S{source_number}",
        "module": lesson,
        "kind": "essay",
        "hintCase": "source",
        "section": "Historical Source Analysis Essays",
        "title": f"{lesson} source {source_number}",
        "slug": f"essay-{lesson}-source-{source_number}",
        "sourceTitle": source_title,
        "items": items,
    }

units = [make_unit(letter) for letter in "BCDEFGH"]
data_text = "// Generated exactly from the model-essay tables in the supplied Module 17 DOCX.\n"
data_text += "window.QUIZ_DATA.units.push(..." + json.dumps(units, ensure_ascii=False, indent=2) + ");\n"
(ROOT / "essay-data-2-8.js").write_text(data_text, encoding="utf-8")

template = """<!doctype html>
<html lang=\"en\">
<head>
  <meta charset=\"utf-8\">
  <meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">
  <meta name=\"theme-color\" content=\"#172554\">
  <title>{title} · Complete Analytical Essay</title>
  <link rel=\"stylesheet\" href=\"styles.css?v=20261013c\">
</head>
<body data-unit=\"{code}\">
  <noscript>This practice needs JavaScript enabled.</noscript>
  <main id=\"app\"></main>
  <script src=\"config.js\"></script>
  <script src=\"essay-data.js?v=20261013c\"></script>
  <script src=\"essay-data-2-8.js?v=20261013c\"></script>
  <script src=\"quiz.js?v=20261013c\"></script>
</body>
</html>
"""
for unit in units:
    (ROOT / f"{unit['slug']}.html").write_text(template.format(title=unit["title"], code=unit["code"]), encoding="utf-8")

print(json.dumps({
    "units": [{"code": unit["code"], "sentences": len(unit["items"]), "blanks": sum(len(item["blankWords"]) for item in unit["items"])} for unit in units],
    "total_sentences": sum(len(unit["items"]) for unit in units),
    "total_blanks": sum(len(item["blankWords"]) for unit in units for item in unit["items"]),
}, ensure_ascii=False, indent=2))

