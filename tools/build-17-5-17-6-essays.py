import json
import re
import sys
from pathlib import Path

from docx import Document


SOURCE_PATH = Path(sys.argv[1])
ROOT = Path(__file__).resolve().parents[1]
document = Document(SOURCE_PATH)

STOPWORDS = set("""
a an and are as at be because been being before but by can could did do does each either even for from
had has have he her here him his how i if in into is it its itself may might more most much must no nor
not now of on one only or other our ours out over she should so some such than that the their theirs them
themselves then there these they this those through to too under up us very was we were what when where
whether which while who whom whose why will with would you your yours first second third however therefore
also already later fairly quite really just still again together today yesterday tomorrow after behind many
notably economically politically socially during across year years time times day days way ways part parts thing things
""".split())

MONTHS = set("january february march april may june july august september october november december".split())

# Names, titles, places, and dates stay visible so the blanks test analytical language.
KNOWN_NAMES = set("""
america american americans asia asian britain british burma cambodia central china chinese christ chulalongkorn
cleveland congress dole dutch eduardo europe european europeans ferry france french furnivall galeano germany
god hawaii hawaiian hawaiians honolulu india indies indochina indonesia james java jules latin liliuokalani
malay malaya malaysia malays mckinley mongkut netherlands orient panama philippines filipino filipinos portugal
raffles roosevelt rusling saigon sanford siam siamese singapore somerset spain stamford states theodore thomas
uruguay united vietnam vietnamese
""".split())

TITLES = set("""
baron emperor empress general governor king khedive lord madam minister president prime prince princes queen
saint sir tsar
""".split())

ALLOWED_CAPITALIZED = set("""
Social Darwinism Industrial Revolution European British French American Dutch Chinese Indian Muslim Hindu
Christian Christianity
""".split())

THEME_WORDS = set("""
academic advanced advancement annex annexation authority biased bias buffer business canal capital capitalist
colonial colonialism colonies colony commercial competition constitution constitutional control credible culture
debt debts direct economic economy education effects empire evidence exploitation exports factories factory
foreign government historical imperial imperialism imperialist imperialists independence independent industries
industry influence information intervention investments labor land lenders limitation loans markets military
modern modernized monopoly motives neutral occupation paternalism perspective political power primary profits
proclamation propaganda raw reliable research resistance rights royal rule ruled ruler secondary slavery social
society source sphere strategic superiority tariff territory trade value wealth
""".split())

ACTION_WORDS = set("""
acted annexed argued became believed borrowed built called changed claimed colonized competed controlled created
develop developped ended encouraged escaped experienced explain exploited exported failed founded gained grew
helped indicates interfered invaded justified lived lost made mentioned modernized occupied persuade plotted
produce promoted protect protested reveals resisted reformed reorganized restore ruled says shows stayed studied
support took traded transformed transmuted wanted warned writes written yielded
""".split())

TOKEN_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ0-9]+(?:['’\-][A-Za-zÀ-ÖØ-öø-ÿ0-9]+)*")


def normalize(value):
    return value.lower().replace("’", "'")


def split_sentences(text):
    protected = re.sub(r"\b(?:Mr|Mrs|Ms|Dr|St|vs|etc)\.", lambda m: m.group(0).replace(".", "∯"), text)
    protected = re.sub(r"\b(?:[A-Z]\.){2,}", lambda m: m.group(0).replace(".", "∯"), protected)
    protected = re.sub(r"\b[A-Z]\.\s*(?=[A-Z][a-z])", lambda m: m.group(0).replace(".", "∯"), protected)
    parts = re.split(r"(?:(?<=[.!?][”’\"'])|(?<=[.!?])|(?<=[.!?]\)))\s+(?=[A-Z0-9“\"'])", protected)
    return [part.replace("∯", ".").strip() for part in parts if part.strip()]


def eligible(token, position):
    raw = token.group(0)
    word = normalize(raw)
    name_base = word[:-2] if word.endswith("'s") else word
    if not word or any(char.isdigit() for char in word) or word in STOPWORDS or word in MONTHS:
        return False
    if len(word) < 3 or name_base in KNOWN_NAMES or word in TITLES:
        return False
    if raw.isupper() and len(raw) <= 4:
        return False
    if raw[0].isupper() and raw not in ALLOWED_CAPITALIZED and position != 0:
        return False
    if position == 0 and name_base in KNOWN_NAMES | TITLES:
        return False
    return True


def word_score(word):
    normalized = normalize(word)
    score = 0
    if normalized in THEME_WORDS:
        score += 7
    if normalized in ACTION_WORDS or normalized.endswith(("ed", "ing", "ize", "ise", "fy")):
        score += 6
    if normalized.endswith(("tion", "sion", "ment", "ness", "ity", "ism", "ship", "ance", "ence")):
        score += 4
    if normalized.endswith(("al", "ic", "ive", "ous", "able", "ible", "ary", "ory")):
        score += 3
    score += min(len(normalized), 9) / 10
    return score


def choose_blank_words(sentence):
    matches = list(TOKEN_RE.finditer(sentence))
    candidates = []
    for position, match in enumerate(matches):
        if eligible(match, position):
            candidates.append((word_score(match.group(0)), position, match.group(0)))
    selected = sorted(sorted(candidates, key=lambda item: (-item[0], item[1]))[:10], key=lambda item: item[1])
    return [item[2] for item in selected]


def purpose_for(section, sentence_number):
    if section == "Introduction":
        labels = [
            "Identify the source, author, audience, and historical setting.",
            "Explain the author’s purpose.",
            "State what the source reveals.",
            "Give the overall limitation or perspective judgement.",
        ]
        return labels[min(sentence_number - 1, len(labels) - 1)]
    if section == "Body 1":
        return "Develop the first content-and-context point with evidence and lesson concepts."
    if section == "Body 2":
        return "Develop the second analytical point and explain its effect on local people."
    if section == "Body 3":
        return "Evaluate the source’s credibility, value, perspective, and limitations."
    return "Restate the argument and explain what other perspectives historians need."


SOURCE_MAP = [
    ("17-5", 1, [6], 5),
    ("17-5", 2, [9], 8),
    ("17-5", 3, [12, 13], 11),
    ("17-5", 4, [16, 17], 15),
    ("17-6", 1, [20], 19),
    ("17-6", 2, [23], 22),
    ("17-6", 3, [26], 25),
    ("17-6", 4, [29], 28),
]


def source_attribution(table_index):
    source_text = document.tables[table_index].cell(0, 0).text.strip()
    lines = [line.strip() for line in source_text.splitlines() if line.strip()]
    return next((line for line in reversed(lines) if line.startswith("—")), lines[-1])


def make_unit(lesson, source_number, table_indexes, source_table_index):
    sections = []
    for table_index in table_indexes:
        for row in document.tables[table_index].rows:
            heading_lines = [line.strip() for line in row.cells[0].text.splitlines() if line.strip()]
            if not heading_lines or heading_lines[0].startswith("Mind map"):
                continue
            section = heading_lines[0].rstrip(":")
            if section.startswith("Body 3"):
                section = "Body 3"
            section_text = row.cells[1].text.strip()
            if section in {"Introduction", "Body 1", "Body 2", "Body 3", "Conclusion"} and section_text:
                sections.append((section, section_text))

    expected = ["Introduction", "Body 1", "Body 2", "Body 3", "Conclusion"]
    if [section for section, _ in sections] != expected:
        raise ValueError(f"Unexpected essay structure for {lesson} source {source_number}: {[s for s, _ in sections]}")

    items = []
    item_id = 1
    for section, section_text in sections:
        sentences = split_sentences(section_text)
        if " ".join(sentences) != section_text:
            raise ValueError(f"Sentence split changed source text for {lesson} source {source_number} {section}")
        for section_sentence, sentence in enumerate(sentences, start=1):
            blank_words = choose_blank_words(sentence)
            if not blank_words:
                raise ValueError(f"No eligible blanks for {lesson} source {source_number} {section} sentence {section_sentence}")
            items.append({
                "id": item_id,
                "essaySection": section,
                "sectionSentence": section_sentence,
                "sectionCount": len(sentences),
                "focus": "Action verbs, theme nouns, and analytical adjectives",
                "term": section,
                "pos": f"Sentence {section_sentence} of {len(sentences)}",
                "definition": purpose_for(section, section_sentence),
                "blankWords": blank_words,
                "prompt": sentence,
            })
            item_id += 1

    return {
        "code": f"{lesson}-S{source_number}",
        "module": lesson,
        "kind": "essay",
        "hintCase": "source",
        "section": f"{lesson} Historical Source Analysis Essays",
        "title": f"{lesson} source {source_number}",
        "slug": f"essay-{lesson}-source-{source_number}",
        "sourceTitle": source_attribution(source_table_index),
        "items": items,
    }


units = [make_unit(*entry) for entry in SOURCE_MAP]
data_text = "// Generated exactly from the model-essay tables in the supplied 17-5/17-6 DOCX.\n"
data_text += "window.QUIZ_DATA.units.push(..." + json.dumps(units, ensure_ascii=False, indent=2) + ");\n"
(ROOT / "essay-data-17-5-17-6.js").write_text(data_text, encoding="utf-8")

template = """<!doctype html>
<html lang=\"en\">
<head>
  <meta charset=\"utf-8\">
  <meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">
  <meta name=\"theme-color\" content=\"#172554\">
  <title>{title} · Complete Analytical Essay</title>
  <link rel=\"stylesheet\" href=\"styles.css?v=20261014a\">
</head>
<body data-unit=\"{code}\">
  <noscript>This practice needs JavaScript enabled.</noscript>
  <main id=\"app\"></main>
  <script src=\"config.js\"></script>
  <script src=\"essay-data.js?v=20261014a\"></script>
  <script src=\"essay-data-2-8.js?v=20261014a\"></script>
  <script src=\"essay-data-17-5-17-6.js?v=20261014a\"></script>
  <script src=\"quiz.js?v=20261014a\"></script>
</body>
</html>
"""
for unit in units:
    (ROOT / f"{unit['slug']}.html").write_text(template.format(title=unit["title"], code=unit["code"]), encoding="utf-8")

print(json.dumps({
    "units": [{
        "code": unit["code"],
        "sentences": len(unit["items"]),
        "blanks": sum(len(item["blankWords"]) for item in unit["items"]),
        "source": unit["sourceTitle"],
    } for unit in units],
    "total_sentences": sum(len(unit["items"]) for unit in units),
    "total_blanks": sum(len(item["blankWords"]) for unit in units for item in unit["items"]),
}, ensure_ascii=False, indent=2))
