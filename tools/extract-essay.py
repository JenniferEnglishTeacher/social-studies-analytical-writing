import json
import re
import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
document = Document(sys.argv[1])
letter = sys.argv[2].upper()
letters = "ABCDEFGH"
if letter not in letters:
    raise SystemExit("Use a source letter from A to H")
essay_table_index = 8 + letters.index(letter) * 6
table = document.tables[essay_table_index]
title = document.tables[essay_table_index - 4].cell(0, 0).text.strip()

def split_sentences(text):
    protected = text
    replacements = {}
    patterns = [
        r"\b(?:Mr|Mrs|Ms|Dr|St|vs|etc)\.",
        r"\b[A-Z]\.\s*(?=[A-Z]\.)",
        r"\b[A-Z]\.\s*(?=[A-Z][a-z])",
    ]
    for pattern in patterns:
        for match in list(re.finditer(pattern, protected)):
            token = f"<DOT{len(replacements)}>"
            replacements[token] = match.group(0).replace(".", "∯")
            protected = protected[:match.start()] + token + protected[match.end():]
            break
    parts = re.split(r"(?:(?<=[.!?][”’\"'])|(?<=[.!?])|(?<=[.!?]\)))\s+(?=[A-Z0-9“\"'])", protected)
    restored = []
    for part in parts:
        for token, value in replacements.items():
            part = part.replace(token, value)
        restored.append(part.replace("∯", ".").strip())
    return [part for part in restored if part]

sections = []
for row in table.rows[1:]:
    heading_lines = [line.strip() for line in row.cells[0].text.splitlines() if line.strip()]
    sections.append({
        "section": heading_lines[0],
        "focus": heading_lines[1] if len(heading_lines) > 1 else "",
        "text": row.cells[1].text.strip(),
        "sentences": split_sentences(row.cells[1].text.strip()),
    })
print(json.dumps({"source": letter, "title": title, "sections": sections}, ensure_ascii=False, indent=2))

